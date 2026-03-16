from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException, status, UploadFile, File, Form
from sqlalchemy.orm import Session
from typing import List, Optional
from app.database import get_db
from app.api.dependencies import get_current_active_user
from app.models.user import User
from app.models.quiz import Quiz
from app.models.question import Question
from app.models.option import Option
from app.models.attempt import Attempt
from app.schemas.quiz import QuizCreate, QuizResponseSchema, QuizSimple
from app.schemas.stats import InstructorStats, LeaderboardEntry
from sqlalchemy import func

from openpyxl import load_workbook

router = APIRouter()

@router.get("/quizzes", response_model=List[QuizSimple])
def get_quizzes(db: Session = Depends(get_db), current_user: User = Depends(get_current_active_user)):
    if current_user.role != "Instructor":
        raise HTTPException(status_code=403, detail="Only instructors can access this")
    
    quizzes = db.query(Quiz).filter(Quiz.instructor_id == current_user.id).all()
    result = []
    for quiz in quizzes:
        q_simple = QuizSimple(
            id=quiz.id,
            title=quiz.title,
            timer=quiz.timer,
            questions_count=len(quiz.questions)
        )
        result.append(q_simple)
    return result

@router.post("/quizzes", response_model=QuizResponseSchema)
def create_quiz(quiz_in: QuizCreate, db: Session = Depends(get_db), current_user: User = Depends(get_current_active_user)):
    if current_user.role != "Instructor":
        raise HTTPException(status_code=403, detail="Only instructors can create quizzes")
    
    db_quiz = Quiz(title=quiz_in.title, timer=quiz_in.timer, instructor_id=current_user.id)
    db.add(db_quiz)
    db.commit()
    db.refresh(db_quiz)

    for q_in in quiz_in.questions:
        db_question = Question(quiz_id=db_quiz.id, question_text=q_in.question_text)
        db.add(db_question)
        db.commit()
        db.refresh(db_question)
        
        for o_in in q_in.options:
            db_option = Option(question_id=db_question.id, option_text=o_in.option_text, is_correct=o_in.is_correct)
            db.add(db_option)
    
    db.commit()
    db.refresh(db_quiz)
    return db_quiz


def _parse_excel_quiz(file_contents: bytes):
    """Parse Excel file into (questions list)."""
    wb = load_workbook(filename=BytesIO(file_contents), read_only=True)
    ws = wb.active

    # Expect header row
    headers = [str(cell.value).strip().lower() if cell.value else '' for cell in next(ws.iter_rows(min_row=1, max_row=1))]

    # Common fields
    question_idx = None
    option_cols = []
    correct_col = None

    for i, h in enumerate(headers):
        if h in ("question", "question_text", "question text"):
            question_idx = i
        elif h in ("option1", "option 1", "option_1", "a"):
            option_cols.append(i)
        elif h in ("option2", "option 2", "option_2", "b"):
            option_cols.append(i)
        elif h in ("option3", "option 3", "option_3", "c"):
            option_cols.append(i)
        elif h in ("option4", "option 4", "option_4", "d"):
            option_cols.append(i)
        elif h in ("correct", "correct_option", "correct option", "answer"):
            correct_col = i

    if question_idx is None or not option_cols or correct_col is None:
        raise ValueError("Excel must include columns: Question, Option1..Option4, Correct (1-4 or A-D or text)")

    questions = []

    for row in ws.iter_rows(min_row=2):
        question_text = (row[question_idx].value or "").strip()
        if not question_text:
            continue

        options = []
        for opt_i in option_cols:
            val = row[opt_i].value
            if val is None:
                options.append("")
            else:
                options.append(str(val).strip())

        # Determine correct index
        correct_raw = row[correct_col].value
        correct_index = None
        if correct_raw is not None:
            raw_str = str(correct_raw).strip()
            if raw_str.isdigit():
                idx = int(raw_str) - 1
                if 0 <= idx < len(options):
                    correct_index = idx
            elif len(raw_str) == 1 and raw_str.upper() in "ABCD":
                idx = ord(raw_str.upper()) - 65
                if 0 <= idx < len(options):
                    correct_index = idx
            else:
                for opt_idx, opt in enumerate(options):
                    if opt and opt.lower() == raw_str.lower():
                        correct_index = opt_idx
                        break

        questions.append({
            "question_text": question_text,
            "options": [{"option_text": opt, "is_correct": idx == correct_index} for idx, opt in enumerate(options) if opt],
        })

    return questions


@router.post("/quizzes/import", response_model=QuizResponseSchema)
def import_quiz(
    title: str = Form(...),
    timer: int = Form(...),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user),
):
    if current_user.role != "Instructor":
        raise HTTPException(status_code=403, detail="Only instructors can import quizzes")

    if file.content_type not in (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/vnd.ms-excel",
    ):
        raise HTTPException(status_code=400, detail="Invalid file type. Upload an Excel (.xlsx) file.")

    try:
        contents = file.file.read()
        questions = _parse_excel_quiz(contents)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to parse Excel file: {e}")

    if not questions:
        raise HTTPException(status_code=400, detail="No questions found in the uploaded file.")

    # Create quiz from parsed questions
    quiz_payload = QuizCreate(title=title, timer=timer, questions=questions)
    return create_quiz(quiz_payload, db=db, current_user=current_user)

@router.get("/quizzes/{id}", response_model=QuizResponseSchema)
def get_quiz(id: int, db: Session = Depends(get_db), current_user: User = Depends(get_current_active_user)):
    quiz = db.query(Quiz).filter(Quiz.id == id).first()
    if not quiz:
        raise HTTPException(status_code=404, detail="Quiz not found")
    return quiz

@router.delete("/quizzes/{id}")
def delete_quiz(id: int, db: Session = Depends(get_db), current_user: User = Depends(get_current_active_user)):
    quiz = db.query(Quiz).filter(Quiz.id == id, Quiz.instructor_id == current_user.id).first()
    if not quiz:
        raise HTTPException(status_code=404, detail="Quiz not found or unauthorized")
    db.delete(quiz)
    db.commit()
    return {"message": "Quiz deleted successfully"}

@router.get("/stats", response_model=InstructorStats)
def get_stats(db: Session = Depends(get_db), current_user: User = Depends(get_current_active_user)):
    if current_user.role != "Instructor":
        raise HTTPException(status_code=403, detail="Unauthorized")
    
    total_quizzes = db.query(Quiz).filter(Quiz.instructor_id == current_user.id).count()
    quiz_ids = [q.id for q in db.query(Quiz.id).filter(Quiz.instructor_id == current_user.id).all()]
    
    total_attempts = db.query(Attempt).filter(Attempt.quiz_id.in_(quiz_ids)).count() if quiz_ids else 0
    avg_score = db.query(func.avg(Attempt.score)).filter(Attempt.quiz_id.in_(quiz_ids)).scalar() if quiz_ids else 0
    active_students = db.query(Attempt.student_id).filter(Attempt.quiz_id.in_(quiz_ids)).distinct().count() if quiz_ids else 0
    
    return {
        "totalQuizzes": total_quizzes,
        "totalAttempts": total_attempts,
        "averageScore": round(avg_score or 0, 2),
        "activeStudents": active_students
    }

@router.get("/leaderboard", response_model=List[LeaderboardEntry])
def get_leaderboard(db: Session = Depends(get_db), current_user: User = Depends(get_current_active_user)):
    # Global leaderboard for simplicity
    results = db.query(
        User.id, 
        User.name, 
        func.sum(Attempt.score).label("total_score"),
        func.count(Attempt.id).label("total_attempts")
    ).join(Attempt, User.id == Attempt.student_id).group_by(User.id).order_by(func.sum(Attempt.score).desc()).all()
    
    leaderboard = []
    for i, res in enumerate(results):
        leaderboard.append({
            "id": res.id,
            "name": res.name,
            "score": res.total_score,
            "attempts": res.total_attempts,
            "rank": i + 1,
            "isCurrentUser": res.id == current_user.id
        })
    return leaderboard

@router.put("/quizzes/{id}", response_model=QuizResponseSchema)
def update_quiz(id: int, quiz_in: QuizCreate, db: Session = Depends(get_db), current_user: User = Depends(get_current_active_user)):
    quiz = db.query(Quiz).filter(Quiz.id == id, Quiz.instructor_id == current_user.id).first()
    if not quiz:
        raise HTTPException(status_code=404, detail="Quiz not found")
    
    quiz.title = quiz_in.title
    quiz.timer = quiz_in.timer
    
    # Clear old questions
    db.query(Question).filter(Question.quiz_id == id).delete()
    db.commit()
    
    for q_in in quiz_in.questions:
        db_question = Question(quiz_id=quiz.id, question_text=q_in.question_text)
        db.add(db_question)
        db.commit()
        db.refresh(db_question)
        
        for o_in in q_in.options:
            db_option = Option(question_id=db_question.id, option_text=o_in.option_text, is_correct=o_in.is_correct)
            db.add(db_option)
    
    db.commit()
    db.refresh(quiz)
    return quiz

@router.get("/analytics")
def get_analytics(db: Session = Depends(get_db), current_user: User = Depends(get_current_active_user)):
    if current_user.role != "Instructor":
        raise HTTPException(status_code=403, detail="Unauthorized")
    
    return {
        "performanceOverTime": [
            { "name": "Week 1", "avg": 65, "high": 90, "low": 40 },
            { "name": "Week 2", "avg": 68, "high": 92, "low": 45 },
            { "name": "Week 3", "avg": 74, "high": 95, "low": 50 },
            { "name": "Week 4", "avg": 81, "high": 98, "low": 60 }
        ],
        "attemptsByQuiz": [
            { "name": "Sample Quiz", "attempts": 5 }
        ]
    }
